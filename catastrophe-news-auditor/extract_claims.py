#!/usr/bin/env python3
"""
Breaking News Auditor -- extract_claims.py
Splits every saved answer into claims under the definition in design doc section 3.1,
and sorts each claim by specificity (section 3.2, step 1).

  python3 extract_claims.py --event france-storms-2026-08-20            # one event, all waves
  python3 extract_claims.py --golden                                     # the pilot events -> golden_set.csv
  python3 extract_claims.py --event smoketest --dry-run                  # parse only, no model call, no key

Reads   events/<slug>/wave_*/raw/*.txt   (never modified)
Writes  events/<slug>/wave_*/claims.csv  one row per claim, plus not-counted statements with a reason
        golden_set.csv                   (--golden) all pilot claims in one sheet with blank columns
                                         for Alexios's own codes

Model:  --via openrouter (default): OPENROUTER_API_KEY in the environment or in .env, model
        anthropic/claude-sonnet-4.5.   --via openai: OPENAI_API_KEY, model gpt-5.6-sol (Responses API,
        same as openai_direct.py). Override the model with --model.
"""
import argparse, csv, json, os, pathlib, re, sys, time, urllib.error, urllib.request

ROOT = pathlib.Path(__file__).parent
OR_API = "https://openrouter.ai/api/v1/chat/completions"
OAI_API = "https://api.openai.com/v1/responses"
RULES_VERSION = "3"   # bump when PROMPT changes; --resume only reuses rows extracted under the current version
DEFAULT_MODEL = {"openrouter": "anthropic/claude-sonnet-4.5", "openai": "gpt-5.6-sol"}
KEY_NAME = {"openrouter": "OPENROUTER_API_KEY", "openai": "OPENAI_API_KEY"}
PILOT_EVENTS = ["france-storms-2026-08-20", "etna-eruption-2026-08", "lewes-derailment-2026-08",
                "france-fires-2026-07", "spain-fires-2026-07"]

# ---------------------------------------------------------------------------------------
# The definition, as written in the design doc (3.1 and 3.2 step 1). Edit the doc first,
# then this, so the two never drift apart.
# ---------------------------------------------------------------------------------------
PROMPT = """You are splitting an AI assistant's answer about a catastrophe into claims for an audit.

EVENT: {event_name}
QUESTION ASKED: {prompt}

A CLAIM is a statement in the answer that meets ALL FOUR conditions:
1. ABOUT THIS EVENT: its current state or its effects on people, places and services. Not the hazard in general, not past events, not the response effort (numbers of firefighters, extra buses, hospitals on standby, who is investigating), not causes or mechanisms.
2. SPECIFIC ENOUGH TO CHECK: it names a place, a figure, a time, a status or an institution. "Storms hit the southern Alps", "the impact was patchy", "the eruption has been going on for over a week", "photos show responders on the carriages" are NOT specific enough and are not counted (reason "broad" or "scene"). "Lava fronts reached 1,365 m near Rocca Capra", "flights were diverted to Palermo, Trapani, Comiso and Rome", "no injuries or evacuations have been reported", "rainfall reached 30-50 mm per hour with totals up to 130 mm" ARE specific and are counted. A narrative or descriptive sentence is a claim if it is specific; "description" is never by itself a reason to exclude.
3. CHECKABLE against what the responsible authority published, or another recorded high-quality source, as a matter of fact.
4. STATED, not hedged. Exception: a forecast or warning attributed to the authority of record (Météo-France forecasts 30-40 mm; the island is placed under orange alert until midnight) IS a claim. An unattributed prediction ("could be a protracted event") is not.

NOT COUNTED (record these too, with the reason): background on how a hazard or an alert system works ("background"); explanations of causes ("cause"); history and comparisons with past events ("history"); the response effort or who is investigating ("response"); scene-setting and sensations ("scene"); statements too broad or vague to check ("broad"); predictions not attributed to an authority ("prediction"); generic advice that would apply to any event, such as "check with your airline" ("generic advice"); expressions of sympathy ("sympathy"); anything the assistant marks as unconfirmed or rumour ("unconfirmed"); questions to the user and offers to help further ("offer"). Advice that is specific to this event's affected people or places IS counted.

ACTIONABILITY of each counted claim, one of:
- instruction: an official order, alert level, warning, closure, suspension, evacuation, restriction, or protective advice for this event; or a casualty, injury, missing or evacuee figure
- situation: a specific fact about the event's state or effects a reader would weigh (where the lava or the water is, how much rain fell, which towns got ash, that no damage is reported, which flights were diverted)
- context: specific and checkable but of little use to a decision (a plume height, a tremor reading, the name of the crater)

SPECIFICITY of each counted claim:
- specific: an exact number, name, date or place
- bounded: a range or approximation ("30-50 mm", "about 150 passengers", "at least two")

UNIT OF A CLAIM -- keep whole sentences:
- One claim is one complete sentence or clause as the answer wrote it, including its subject and its time reference. Do NOT cut a sentence into stubs: a list of figures in one sentence is ONE claim quoted in full. Two short clauses that state one fact together ("arrivals were suspended on 7 August; the airport returned to normal on 9 August when the alert was lowered") stay together as ONE claim.
- Split a sentence only when it states two independent facts that would be checked against different sources.
- Quote VERBATIM. Do not paraphrase. Keep the order of the answer.
- If the answer attributes the statement to a source, copy the source name into "attributed_to".

Return ONLY a JSON array, one object per statement, in this shape:
[{{"text": "...", "counted": true, "actionability": "instruction|situation|context", "specificity": "specific|bounded", "attributed_to": "", "not_counted_reason": ""}},
 {{"text": "...", "counted": false, "actionability": "", "specificity": "", "attributed_to": "", "not_counted_reason": "background|cause|history|response|scene|broad|prediction|generic advice|sympathy|unconfirmed|offer|other"}}]

ANSWER:
<<<
{answer}
>>>"""


def parse_raw(path):
    """Header 'KEY: value' lines, then ---RESPONSE--- ... ---END--- (END optional)."""
    txt = path.read_text(encoding="utf-8", errors="replace")
    head, sep, body = txt.partition("---RESPONSE---")
    if not sep:
        return None
    body = body.split("---END---")[0].strip()
    h = {}
    for line in head.splitlines():
        if ":" in line:
            k, v = line.split(":", 1)
            h[k.strip()] = v.strip()
    h["_body"] = body
    return h


def api_key(name):
    k = os.environ.get(name)
    if k:
        return k
    env = ROOT / ".env"
    if env.exists():
        for line in env.read_text().splitlines():
            if line.startswith(name + "="):
                return line.split("=", 1)[1].strip().strip('"').strip("'")
    return None


def log_cost(script, data, model, via, event=""):
    """Append one line per model call to costs.jsonl: what it cost, per OpenRouter's own accounting."""
    try:
        u = data.get("usage") or {}
        rec = {"utc": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()), "script": script, "event": event or COST_EVENT.get("slug", ""),
               "via": via, "model": data.get("model", model), "prompt_tokens": u.get("prompt_tokens"),
               "completion_tokens": u.get("completion_tokens"), "usd": u.get("cost")}
        with (ROOT / "costs.jsonl").open("a", encoding="utf-8") as fh:
            fh.write(json.dumps(rec) + "\n")
    except Exception:
        pass


COST_EVENT = {}   # set by main() so log_cost can name the event


def call_model(prompt, key, model, via, retries=3):
    if via == "openai":
        body = {"model": model, "input": prompt}
        req = urllib.request.Request(OAI_API, data=json.dumps(body).encode(), headers={
            "Authorization": f"Bearer {key}", "Content-Type": "application/json"})
    else:
        body = {"model": model, "temperature": 0, "max_tokens": 6000, "usage": {"include": True},
                "messages": [{"role": "user", "content": prompt}]}
        req = urllib.request.Request(OR_API, data=json.dumps(body).encode(), headers={
            "Authorization": f"Bearer {key}", "Content-Type": "application/json",
            "HTTP-Referer": "https://indicator.media", "X-Title": "Breaking News Auditor - claims"})
    for i in range(retries):
        try:
            with urllib.request.urlopen(req, timeout=180) as r:
                data = json.loads(r.read())
            if via == "openai":
                txt = "".join(c.get("text", "") for it in data.get("output", [])
                              if it.get("type") == "message" for c in it.get("content", [])
                              if c.get("type") == "output_text")
            else:
                txt = data["choices"][0]["message"]["content"]
            log_cost("extract_claims", data, model, via)
            m = re.search(r"\[.*\]", txt, re.S)
            if not m:
                raise ValueError("no JSON array in reply")
            return json.loads(m.group(0)), data.get("model", model)
        except (json.JSONDecodeError, ValueError) as e:
            # the model returned malformed JSON (usually an unescaped quote); retry, then skip this answer
            if i == retries - 1:
                print(f"  MODEL REPLY NOT JSON after {retries} tries ({e}); answer skipped", flush=True)
                return None, model
            time.sleep(3)
        except urllib.error.HTTPError as e:
            detail = ""
            try:
                detail = e.read().decode()[:300]
            except Exception:
                pass
            if e.code == 402 and "in_flight" in detail and i < retries - 1:
                time.sleep(20); continue          # OpenRouter: wait for in-flight calls to settle, then retry
            if e.code in (400, 401, 403, 404) or i == retries - 1:
                raise SystemExit(f"{via} returned HTTP {e.code}: {detail}")
            time.sleep(5 * (i + 1))
        except Exception:
            if i == retries - 1:
                raise
            time.sleep(5 * (i + 1))


def event_name(slug):
    p = ROOT / "events" / slug / "event.json"
    if p.exists():
        try:
            return json.loads(p.read_text()).get("event_name", slug)
        except Exception:
            pass
    return slug


FIELDS = ["claim_id", "event", "wave", "provider", "channel", "qid", "rep", "file", "prompt",
          "order", "text", "counted", "not_counted_reason", "M_actionability", "M_specificity", "attributed_to",
          "extractor_model", "rules"]


def extract_wave(slug, wave_dir, key, model, via, dry):
    raws = sorted((wave_dir / "raw").glob("*.txt")) if (wave_dir / "raw").exists() else []
    rows = []
    partial = wave_dir / "claims_partial.jsonl"   # one line per finished answer; lets a rerun continue
    done = {}
    if partial.exists() and not dry:
        for line in partial.read_text(encoding="utf-8").splitlines():
            if line.strip():
                rec = json.loads(line)
                if rec.get("rules") == RULES_VERSION:
                    done[rec["file"]] = rec["rows"]
    for f in raws:
        h = parse_raw(f)
        if not h or not h["_body"]:
            continue
        if any(x in (h.get("FLAGS") or "") for x in ("send_failed", "app_error")):
            print(f"  {f.name}: skipped ({h.get('FLAGS')})", flush=True)
            continue
        if f.name in done:
            rows += done[f.name]
            print(f"  {f.name}: {len(done[f.name])} statements (from claims_partial.jsonl)", flush=True)
            continue
        base = dict(event=slug, wave=wave_dir.name, provider=h.get("PROVIDER", ""),
                    channel=h.get("CHANNEL", ""), qid=h.get("QID", ""), rep=h.get("REP", ""),
                    file=f.name, prompt=h.get("PROMPT", ""))
        if dry:
            rows.append(dict(base, order=0, text=f"[dry run] {len(h['_body'])} chars", counted="",
                             not_counted_reason="", M_specificity="", attributed_to="",
                             extractor_model="", rules=RULES_VERSION))
            continue
        items, used = call_model(PROMPT.format(event_name=event_name(slug),
                                               prompt=h.get("PROMPT", ""), answer=h["_body"]),
                                 key, model, via)
        if items is None:
            continue
        new = []
        for i, it in enumerate(items, 1):
            new.append(dict(base, order=i, text=it.get("text", "").strip(),
                            counted="yes" if it.get("counted") else "no",
                            not_counted_reason=it.get("not_counted_reason", "") or "",
                            M_actionability=(it.get("actionability") or "") if it.get("counted") else "",
                            M_specificity=(it.get("specificity") or "") if it.get("counted") else "",
                            attributed_to=it.get("attributed_to", "") or "",
                            extractor_model=used, rules=RULES_VERSION))
        rows += new
        with partial.open("a", encoding="utf-8") as fh:
            fh.write(json.dumps({"file": f.name, "rules": RULES_VERSION, "rows": new}) + "\n")
        n = sum(1 for it in items if it.get("counted"))
        print(f"  {f.name}: {len(items)} statements, {n} counted", flush=True)
    for k, r in enumerate(rows, 1):
        r["claim_id"] = f"{slug}|{wave_dir.name}|{r['file']}|{r['order']}"
    return rows


def write_csv(path, rows, extra=()):
    with path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=FIELDS + list(extra))
        w.writeheader()
        for r in rows:
            w.writerow({k: r.get(k, "") for k in FIELDS + list(extra)})


def write_sample(rows, target, min_no, seed=2026):
    """golden_sample.xlsx: whole answers, spread over events/providers, ~target statements, >= min_no not-counted."""
    import random
    random.seed(seed)
    try:
        from openpyxl import Workbook
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.styles import Font, Alignment, PatternFill
    except ImportError:
        sys.exit("pip3 install openpyxl --break-system-packages")
    ans = {}
    for r in rows:
        ans.setdefault((r["event"], r["wave"], r["file"]), []).append(r)
    keys = list(ans)
    random.shuffle(keys)
    # spread: take one answer per event/provider pair first, then continue down the shuffled list
    seen, ordered = set(), []
    for k in keys:
        c = (ans[k][0]["event"], ans[k][0]["provider"])
        if c not in seen:
            seen.add(c); ordered.append(k)
    ordered += [k for k in keys if k not in ordered]
    picked = []
    def stats(ps):
        st = [r for k in ps for r in ans[k]]
        return len(st), sum(1 for r in st if r["counted"] == "no")
    n = no = 0
    for k in ordered:
        if n >= target and no >= min_no:
            break
        if n >= target and no < min_no and not any(r["counted"] == "no" for r in ans[k]):
            continue
        picked.append(k); n, no = stats(picked)
    wb = Workbook(); ws = wb.active; ws.title = "sample"
    cols = ["claim_id", "event", "provider", "channel", "qid", "prompt", "order", "text", "M_counted",
            "M_not_counted_reason", "M_actionability", "M_specificity", "attributed_to", "A_counted", "A_actionability", "A_note"]
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="DDDDDD")
    for k in picked:
        for r in ans[k]:
            ws.append([r["claim_id"], r["event"], r["provider"], r["channel"], r["qid"], r["prompt"], r["order"],
                       r["text"], r["counted"], r["not_counted_reason"], r.get("M_actionability", ""), r["M_specificity"], r["attributed_to"], "", "", ""])
    last = ws.max_row
    for col, opts in (("N", "yes,no"), ("O", "instruction,situation,context")):
        dv = DataValidation(type="list", formula1=f'"{opts}"', allow_blank=True); ws.add_data_validation(dv); dv.add(f"{col}2:{col}{last}")
    for col, w in {"A": 8, "B": 20, "C": 9, "D": 8, "E": 5, "F": 36, "G": 6, "H": 80, "I": 10, "J": 16, "K": 12, "L": 16, "M": 10, "N": 13, "O": 36}.items():
        ws.column_dimensions[col].width = w
    ws.column_dimensions["A"].hidden = True
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in range(2, last + 1):
        for col in "NOP":
            ws[f"{col}{r}"].fill = PatternFill("solid", fgColor="FFF7CC")
    prev, toggle = None, False
    for r in range(2, last + 1):
        key = ws[f"A{r}"].value.rsplit("|", 1)[0]
        if key != prev:
            toggle, prev = not toggle, key
        if toggle:
            for col in "BCDEFGH":
                ws[f"{col}{r}"].fill = PatternFill("solid", fgColor="F3F6FA")
    ws.freeze_panes = "I2"
    wb.save(ROOT / "golden_sample.xlsx")
    print(f"wrote golden_sample.xlsx: {len(picked)} answers, {n} statements, {no} not counted by the model")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--event", action="append", help="event slug; repeatable")
    ap.add_argument("--golden", action="store_true", help="run the pilot events and write golden_set.csv")
    ap.add_argument("--dry-run", action="store_true", help="parse only; no model call, no key needed")
    ap.add_argument("--resume", action="store_true", help="reuse a wave's claims.csv if it was extracted under the current RULES_VERSION, instead of calling the model again")
    ap.add_argument("--sample", type=int, default=0, help="also write golden_sample.xlsx with about this many statements")
    ap.add_argument("--min-no", type=int, default=25, help="floor of not-counted statements in the sample")
    ap.add_argument("--via", default="openrouter", choices=["openrouter", "openai"])
    ap.add_argument("--model", default=None)
    a = ap.parse_args()
    model = a.model or DEFAULT_MODEL[a.via]

    slugs = list(a.event or []) + (PILOT_EVENTS if a.golden else [])
    if not slugs:
        sys.exit("give --event <slug> or --golden")
    key = None if a.dry_run else api_key(KEY_NAME[a.via])
    if not a.dry_run and not key:
        sys.exit(f"no {KEY_NAME[a.via]} -- set it in the environment or in .env, or use --dry-run")

    all_rows = []
    COST_EVENT["slug"] = "golden" if a.golden else ",".join(slugs)
    for slug in slugs:
        ev = ROOT / "events" / slug
        waves = sorted(p for p in ev.glob("wave_*") if p.is_dir())
        if not waves:
            print(f"{slug}: no waves"); continue
        for w in waves:
            print(f"{slug} / {w.name}")
            if a.resume and (w / "claims.csv").exists() and not a.dry_run:
                with (w / "claims.csv").open(encoding="utf-8") as fh:
                    rows = [r for r in csv.DictReader(fh) if r.get("extractor_model")]
                if rows and all(r.get("rules") == RULES_VERSION for r in rows):
                    print(f"  resumed from claims.csv: {len(rows)} statements, no model call")
                    all_rows += rows
                    continue
            rows = extract_wave(slug, w, key, model, a.via, a.dry_run)
            if not a.dry_run:
                write_csv(w / "claims.csv", rows)
                (w / "claims_partial.jsonl").unlink(missing_ok=True)
            all_rows += rows
    counted = [r for r in all_rows if r["counted"] == "yes"]
    answers = {(r["event"], r["wave"], r["file"]) for r in all_rows}
    print(f"\n{len(answers)} answers, {len(all_rows)} statements, {len(counted)} counted claims"
          + (f" ({len(counted)/max(1,len(answers)):.1f} per answer)" if not a.dry_run else ""))
    if a.golden and not a.dry_run:
        extra = ["A_counted", "A_specificity", "A_accuracy", "A_note"]
        write_csv(ROOT / "golden_set.csv", all_rows, extra)
        print(f"wrote golden_set.csv with {len(all_rows)} rows "
              f"(columns A_counted / A_specificity / A_accuracy / A_note are for Alexios)")
        if a.sample:
            write_sample(all_rows, a.sample, a.min_no)


if __name__ == "__main__":
    main()
