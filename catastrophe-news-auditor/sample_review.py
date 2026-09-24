#!/usr/bin/env python3
"""
Breaking News Auditor -- sample_review.py   (design doc 3.5 and 3.6)
Two jobs:
  1. draw the 10% human-review sample of consumer-app claims for an event, stratified by assistant and wave
       python3 sample_review.py sample --event kathmandu-flood-2026-09-22 [--share 0.10]
     -> events/<slug>/review_sample.xlsx  (claims with the model's codes and blank columns for Alexios)
  2. compute agreement between Alexios's codes and the model's, on any sheet that has both
       python3 sample_review.py agree --file golden_sample.xlsx            # golden set (columns M_* and A_*)
       python3 sample_review.py agree --file events/<slug>/review_sample.xlsx
       python3 sample_review.py agree --file ratings.csv --model-col M_counted --human-col A_counted
     -> raw agreement, Krippendorff's alpha (nominal), confusion table, per code; rows with a blank human code are skipped

Krippendorff's alpha is computed for two coders, nominal data, from the coincidence matrix
(Krippendorff 2011, "Computing Krippendorff's Alpha-Reliability"). 1.0 = perfect, 0 = chance.
"""
import argparse, csv, pathlib, random, sys, collections

ROOT = pathlib.Path(__file__).parent
PAIRS = [("M_counted", "A_counted"), ("M_specificity", "A_specificity"), ("M_accuracy", "A_accuracy"),
         ("M_bounded", "A_bounded"), ("M_time_anchored", "A_time_anchored")]


def read_table(path):
    path = pathlib.Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True).worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(it)]
        return [dict(zip(header, ["" if v is None else str(v).strip() for v in row])) for row in it]
    with path.open(encoding="utf-8") as f:
        return list(csv.DictReader(f))


def krippendorff_alpha(pairs):
    """pairs: list of (code_a, code_b) for the same unit, nominal. Two coders, no missing values."""
    vals = sorted({a for a, _ in pairs} | {b for _, b in pairs})
    idx = {v: i for i, v in enumerate(vals)}
    k = len(vals)
    if k < 2 or len(pairs) < 2:
        return float("nan")
    # coincidence matrix: each unit contributes to (a,b) and (b,a) with weight 1/(m-1) = 1 for m=2
    o = [[0.0] * k for _ in range(k)]
    for a, b in pairs:
        o[idx[a]][idx[b]] += 1
        o[idx[b]][idx[a]] += 1
    n = sum(map(sum, o))
    nc = [sum(row) for row in o]
    do = sum(o[c][kk] for c in range(k) for kk in range(k) if c != kk)
    de = sum(nc[c] * nc[kk] for c in range(k) for kk in range(k) if c != kk) / (n - 1)
    return 1 - do / de if de else float("nan")


def agree(rows, mcol, hcol, label):
    pairs = [(r.get(mcol, "").strip().lower(), r.get(hcol, "").strip().lower()) for r in rows
             if r.get(hcol, "").strip() and r.get(mcol, "").strip()]
    if not pairs:
        return None
    raw = sum(1 for a, b in pairs if a == b) / len(pairs)
    alpha = krippendorff_alpha(pairs)
    conf = collections.Counter(pairs)
    print(f"\n{label}: {len(pairs)} rated rows  raw agreement {raw:.1%}  Krippendorff's alpha {alpha:.2f}"
          f"  -> {'ACCEPTED (>=90%)' if raw >= 0.9 else 'below 90%'}")
    codes = sorted({a for a, _ in pairs} | {b for _, b in pairs})
    print("  model \\ human " + " ".join(f"{c[:12]:>12}" for c in codes))
    for a in codes:
        print(f"  {a[:14]:<14} " + " ".join(f"{conf.get((a, b), 0):>12}" for b in codes))
    return raw, alpha, len(pairs)


def cmd_agree(a):
    rows = read_table(a.file)
    if a.model_col and a.human_col:
        agree(rows, a.model_col, a.human_col, f"{a.model_col} vs {a.human_col}")
        return
    any_ = False
    for m, h in PAIRS:
        if m in rows[0] and h in rows[0]:
            # for accuracy-type pairs only compare rows the human counted as claims
            sub = rows if m == "M_counted" else [r for r in rows if r.get("A_counted", "").lower() in ("", "yes")]
            if agree(sub, m, h, f"{m} vs {h}"):
                any_ = True
    if not any_:
        sys.exit("no M_*/A_* column pairs with values found; pass --model-col and --human-col")


def cmd_sample(a):
    ev = ROOT / "events" / a.event
    claims = []
    for w in sorted(ev.glob("wave_*")):
        p = w / "claims.csv"
        if not p.exists():
            continue
        for r in read_table(p):
            if r.get("counted") == "yes" and r.get("channel", "").startswith("ui"):
                r["wave"] = w.name; claims.append(r)
    if not claims:
        sys.exit(f"no consumer-app claims under events/{a.event}/wave_*/claims.csv (run extract_claims.py first)")
    random.seed(a.seed)
    strata = collections.defaultdict(list)
    for r in claims:
        strata[(r["provider"], r["wave"])].append(r)
    picked = []
    for key, rs in sorted(strata.items()):
        n = max(1, round(len(rs) * a.share))
        picked += random.sample(rs, min(n, len(rs)))
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook(); ws = wb.active; ws.title = "review"
    cols = ["claim_id", "provider", "wave", "qid", "rep", "prompt", "text", "M_counted", "M_accuracy", "M_bounded",
            "M_time_anchored", "A_counted", "A_accuracy", "A_bounded", "A_time_anchored", "A_note"]
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="DDDDDD")
    for r in picked:
        ws.append([r.get("claim_id"), r.get("provider"), r.get("wave"), r.get("qid"), r.get("rep"), r.get("prompt"),
                   r.get("text"), r.get("counted"), r.get("accuracy", ""), r.get("bounded", ""), r.get("time_anchored", ""),
                   "", "", "", "", ""])
    last = ws.max_row
    for col, opts in (("L", "yes,no"), ("M", "correct,incorrect,stale,misattributed,unverifiable"), ("N", "yes,no"), ("O", "yes,no")):
        dv = DataValidation(type="list", formula1=f'"{opts}"', allow_blank=True); ws.add_data_validation(dv); dv.add(f"{col}2:{col}{last}")
    for col, w in {"A": 8, "F": 34, "G": 80, "P": 30}.items():
        ws.column_dimensions[col].width = w
    ws.column_dimensions["A"].hidden = True
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for c in row[11:]:
            c.fill = PatternFill("solid", fgColor="FFF7CC")
    ws.freeze_panes = "H2"
    out = ev / "review_sample.xlsx"
    wb.save(out)
    print(f"{len(claims)} consumer-app claims, {len(picked)} sampled ({a.share:.0%}, stratified by assistant x wave) -> {out.relative_to(ROOT)}")


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("sample"); s.add_argument("--event", required=True); s.add_argument("--share", type=float, default=0.10); s.add_argument("--seed", type=int, default=2026)
    g = sub.add_parser("agree"); g.add_argument("--file", required=True); g.add_argument("--model-col"); g.add_argument("--human-col")
    a = ap.parse_args()
    (cmd_sample if a.cmd == "sample" else cmd_agree)(a)


if __name__ == "__main__":
    main()
